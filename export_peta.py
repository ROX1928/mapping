"""Render detailed, printable area and district maps with store routes."""

import io
import math
import time
import textwrap
import zipfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
import requests
from matplotlib.patches import Circle

AREAS = {
    "A1": ("Karawang Kota", "#0b8a7a", ["Karawang Bar."]),
    "A2": ("Karawang Timur", "#e8710a", ["Klari", "Cikampek"]),
    "A3": ("Pantura Cilamaya-Pamanukan", "#8430ce", ["Cilamaya Wetan", "Pamanukan", "Binong"]),
    "A4": ("Purwakarta", "#1a73e8", ["Purwakarta", "Plered"]),
    "A5": ("Subang Kota", "#d93025", ["Subang"]),
}
KEC2AREA = {district: code for code, (_, _, districts) in AREAS.items() for district in districts}
OSRM_URL = "https://router.project-osrm.org/route/v1/driving/"
OSM_TILE_URL = "https://tile.openstreetmap.org/{z}/{x}/{y}.png"
WEB_MERCATOR_RADIUS = 6_378_137
WEB_MERCATOR_ORIGIN = math.pi * WEB_MERCATOR_RADIUS
TILE_SIZE_METERS = 2 * WEB_MERCATOR_ORIGIN / 256
NEARBY_COLORS = [
    "#d93025", "#188038", "#f29900", "#8430ce", "#00838f", "#c2185b",
    "#1a73e8", "#7b5e00", "#5f6368", "#b3145a", "#137333", "#e37400",
]


def _web_mercator(latitude, longitude):
    latitude = max(-85.05112878, min(85.05112878, float(latitude)))
    x = WEB_MERCATOR_RADIUS * math.radians(float(longitude))
    y = WEB_MERCATOR_RADIUS * math.log(math.tan(math.pi / 4 + math.radians(latitude) / 2))
    return x, y


def _order_stores(group):
    remaining = list(group.index)
    start = min(remaining, key=lambda index: group.at[index, "Longitude"])
    order = [start]
    remaining.remove(start)
    while remaining:
        latitude, longitude = group.loc[order[-1], ["Latitude", "Longitude"]]
        next_index = min(
            remaining,
            key=lambda index: math.hypot(
                group.at[index, "Latitude"] - latitude,
                (group.at[index, "Longitude"] - longitude) * math.cos(math.radians(latitude)),
            ),
        )
        order.append(next_index)
        remaining.remove(next_index)
    return group.loc[order].reset_index(drop=True)


def _interpolate_midpoint(coordinates, fraction):
    points = [(float(lon), float(lat)) for lon, lat in coordinates]
    if len(points) < 2:
        return points[0][1], points[0][0]

    def distance(first, second):
        lat1, lat2 = math.radians(first[1]), math.radians(second[1])
        dlat = lat2 - lat1
        dlon = math.radians(second[0] - first[0])
        h = math.sin(dlat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2) ** 2
        return 12_742_000 * math.asin(min(1, math.sqrt(h)))

    lengths = [distance(first, second) for first, second in zip(points, points[1:])]
    target = sum(lengths) * fraction
    travelled = 0
    for first, second, length in zip(points, points[1:], lengths):
        if length and travelled + length >= target:
            part = (target - travelled) / length
            return (
                first[1] + (second[1] - first[1]) * part,
                first[0] + (second[0] - first[0]) * part,
            )
        travelled += length
    return points[-1][1], points[-1][0]


def _route_for(ordered, title):
    if len(ordered) < 2:
        return [], [], []
    coordinates = ";".join(
        f"{row.Longitude},{row.Latitude}" for row in ordered.itertuples()
    )
    response = requests.get(
        OSRM_URL + coordinates,
        params={"overview": "full", "geometries": "geojson", "steps": "true"},
        timeout=30,
    )
    response.raise_for_status()
    result = response.json()
    routes = result.get("routes", [])
    if not routes:
        raise ValueError(f"Layanan rute tidak memberikan jalur untuk {title}.")
    legs = routes[0].get("legs", [])
    if len(legs) != len(ordered) - 1:
        raise ValueError(f"Data jarak antar toko tidak lengkap untuk {title}.")

    distances = [float(leg["distance"]) / 1000 for leg in legs]
    midpoints = []
    for leg in legs:
        halfway = float(leg["distance"]) / 2
        walked = 0
        midpoint = None
        for step in leg.get("steps", []):
            step_distance = float(step.get("distance", 0))
            geometry = step.get("geometry", {}).get("coordinates", [])
            if not step_distance or not geometry:
                continue
            if walked + step_distance >= halfway:
                midpoint = _interpolate_midpoint(geometry, (halfway - walked) / step_distance)
                break
            walked += step_distance
        if midpoint is None:
            midpoint = (
                (float(ordered.iloc[len(midpoints)]["Latitude"])
                 + float(ordered.iloc[len(midpoints) + 1]["Latitude"])) / 2,
                (float(ordered.iloc[len(midpoints)]["Longitude"])
                 + float(ordered.iloc[len(midpoints) + 1]["Longitude"])) / 2,
            )
        midpoints.append(midpoint)

    route = [(float(lat), float(lon)) for lon, lat in routes[0]["geometry"]["coordinates"]]
    return route, distances, midpoints


def _nearby_colors(ordered):
    positions = [
        _web_mercator(row.Latitude, row.Longitude)
        for row in ordered.itertuples()
    ]
    neighbors = {index: set() for index in range(len(positions))}
    for index, (x1, y1) in enumerate(positions):
        for other in range(index + 1, len(positions)):
            x2, y2 = positions[other]
            if math.hypot(x2 - x1, y2 - y1) <= 300:
                neighbors[index].add(other)
                neighbors[other].add(index)
    colored = {}
    order = sorted(
        (index for index, nearby in neighbors.items() if nearby),
        key=lambda index: (-len(neighbors[index]), index),
    )
    for index in order:
        used = {colored[other] for other in neighbors[index] if other in colored}
        colored[index] = next(
            (color for color in NEARBY_COLORS if color not in used),
            NEARBY_COLORS[len(colored) % len(NEARBY_COLORS)],
        )
    return colored


def _extent(points, axis_width, axis_height):
    xs, ys = zip(*points)
    min_x, max_x = min(xs), max(xs)
    min_y, max_y = min(ys), max(ys)
    width = max(max_x - min_x, 3500)
    height = max(max_y - min_y, 3500)
    center_x, center_y = (min_x + max_x) / 2, (min_y + max_y) / 2
    width *= 1.3
    height *= 1.3
    target_aspect = axis_width / axis_height
    if width / height < target_aspect:
        width = height * target_aspect
    else:
        height = width / target_aspect
    return (
        center_x - width / 2,
        center_x + width / 2,
        center_y - height / 2,
        center_y + height / 2,
    )


def _add_basemap(ax, extent, tile_cache, session):
    min_x, max_x, min_y, max_y = extent
    axis_width = 15 * 100 * 0.70
    axis_height = 9 * 100 * 0.80
    meters_per_pixel = max((max_x - min_x) / axis_width, (max_y - min_y) / axis_height)
    zoom = max(10, min(16, int(math.floor(math.log2(TILE_SIZE_METERS / meters_per_pixel)))))
    tile_span = 2 * WEB_MERCATOR_ORIGIN / (2 ** zoom)
    tile_left = max(0, int(math.floor((min_x + WEB_MERCATOR_ORIGIN) / tile_span)))
    tile_right = min(2 ** zoom - 1, int(math.floor((max_x + WEB_MERCATOR_ORIGIN) / tile_span)))
    tile_top = max(0, int(math.floor((WEB_MERCATOR_ORIGIN - max_y) / tile_span)))
    tile_bottom = min(2 ** zoom - 1, int(math.floor((WEB_MERCATOR_ORIGIN - min_y) / tile_span)))

    for tile_x in range(tile_left, tile_right + 1):
        for tile_y in range(tile_top, tile_bottom + 1):
            key = (zoom, tile_x, tile_y)
            if key not in tile_cache:
                response = session.get(
                    OSM_TILE_URL.format(z=zoom, x=tile_x, y=tile_y),
                    timeout=20,
                )
                response.raise_for_status()
                from PIL import Image

                tile_cache[key] = Image.open(io.BytesIO(response.content)).convert("RGB")
                time.sleep(0.55)
            left = tile_x * tile_span - WEB_MERCATOR_ORIGIN
            right = left + tile_span
            top = WEB_MERCATOR_ORIGIN - tile_y * tile_span
            bottom = top - tile_span
            ax.imshow(
                tile_cache[key],
                extent=(left, right, bottom, top),
                interpolation="bilinear",
                zorder=0,
            )
    ax.set_xlim(min_x, max_x)
    ax.set_ylim(min_y, max_y)
    ax.set_aspect("equal", adjustable="box")
    ax.set_axis_off()


def _render_map(ordered, title, subtitle, area_color, tile_cache, session):
    route, distances, midpoints = _route_for(ordered, title)
    store_points = [
        _web_mercator(row.Latitude, row.Longitude)
        for row in ordered.itertuples()
    ]
    route_points = [_web_mercator(latitude, longitude) for latitude, longitude in route]

    figure = plt.figure(figsize=(15, 9), facecolor="white")
    map_axis = figure.add_axes([0.035, 0.11, 0.68, 0.79])
    table_axis = figure.add_axes([0.735, 0.15, 0.245, 0.72])
    table_axis.set_axis_off()
    figure.suptitle(title, x=0.04, y=0.965, ha="left", fontsize=20, fontweight="bold")
    figure.text(0.04, 0.925, subtitle, ha="left", fontsize=11, color="#4b5563")

    extent = _extent(store_points, 15 * 100 * 0.68, 9 * 100 * 0.79)
    _add_basemap(map_axis, extent, tile_cache, session)
    if route_points:
        map_axis.plot(
            [point[0] for point in route_points],
            [point[1] for point in route_points],
            color=area_color,
            linewidth=4,
            alpha=0.88,
            zorder=3,
        )

    alternate_colors = _nearby_colors(ordered)
    for index, ((x, y), row) in enumerate(zip(store_points, ordered.itertuples())):
        color = alternate_colors.get(index, area_color)
        map_axis.add_patch(
            Circle((x, y), 115, facecolor=color, edgecolor="white", linewidth=2.2, zorder=5)
        )
        map_axis.text(
            x,
            y,
            str(index + 1),
            ha="center",
            va="center",
            fontsize=8,
            fontweight="bold",
            color="white",
            zorder=6,
        )

    for (latitude, longitude), distance in zip(midpoints, distances):
        x, y = _web_mercator(latitude, longitude)
        map_axis.annotate(
            f"{distance:.1f} km",
            (x, y),
            xytext=(0, 11),
            textcoords="offset points",
            ha="center",
            va="bottom",
            fontsize=8,
            fontweight="bold",
            color="#202124",
            bbox={"boxstyle": "round,pad=0.28", "fc": "white", "ec": area_color, "lw": 1.3, "alpha": 0.97},
            zorder=7,
        )

    rows = []
    for index, row in ordered.iterrows():
        distance = "Titik awal" if index == 0 else f"{distances[index - 1]:.1f} km"
        rows.append([
            str(index + 1),
            textwrap.fill(str(row["Nama Toko"]), width=22),
            textwrap.fill(str(row["Kecamatan"]), width=16),
            distance,
        ])
    table_height = min(0.92, max(0.24, 0.06 * (len(rows) + 1)))
    table = table_axis.table(
        cellText=rows,
        colLabels=["#", "Nama toko", "Kecamatan", "Jarak (km)"],
        colWidths=[0.07, 0.37, 0.30, 0.26],
        cellLoc="left",
        colLoc="left",
        bbox=[0, 0.09, 1, table_height],
    )
    table.auto_set_font_size(False)
    table.set_fontsize(8.2 if len(rows) <= 10 else 7.1)
    for (row_index, column), cell in table.get_celld().items():
        cell.set_edgecolor("#d6dbe1")
        cell.set_linewidth(0.6)
        cell.PAD = 0.08
        if row_index == 0:
            cell.set_facecolor("#eaf1fb")
            cell.set_text_props(weight="bold", color="#202124")
        else:
            cell.set_facecolor("#ffffff" if row_index % 2 else "#f7f9fc")
            cell.get_text().set_wrap(True)
            if column == 0:
                cell.set_text_props(weight="bold", color=alternate_colors.get(row_index - 1, area_color))
            elif column == 3:
                cell.set_text_props(weight="bold", color="#202124")

    total_distance = sum(distances)
    summary = f"{len(ordered)} toko · total jalur {total_distance:.1f} km"
    table_axis.text(
        0,
        0.015,
        summary,
        transform=table_axis.transAxes,
        fontsize=9,
        fontweight="bold",
        color="#202124",
    )
    figure.text(
        0.04,
        0.035,
        "Angka pada peta cocok dengan kolom #. Jarak dari toko sebelumnya mengikuti jalan. "
        "Peta © OpenStreetMap contributors.",
        fontsize=8,
        color="#5f6368",
    )
    output = io.BytesIO()
    figure.savefig(output, format="png", dpi=180, facecolor="white")
    plt.close(figure)
    return output.getvalue()


def create_screenshot_images(data):
    required = {"Nama Toko", "Latitude", "Longitude", "Kode Area", "Area", "Kecamatan", "Warna"}
    missing = required - set(data.columns)
    if missing:
        raise ValueError(f"Kolom data toko tidak ditemukan: {', '.join(sorted(missing))}.")
    if data.empty:
        raise ValueError("Tidak ada toko untuk dibuatkan screenshot.")
    if data[["Latitude", "Longitude"]].isna().any().any():
        raise ValueError("Koordinat Latitude/Longitude kosong; lengkapi data sebelum membuat peta.")

    images = {}
    tile_cache = {}
    session = requests.Session()
    session.headers.update({"User-Agent": "MotorolaStoreAreaMap/1.0 (detailed map screenshot export)"})
    try:
        for code, (area_name, area_color, _) in AREAS.items():
            area_group = data[data["Kode Area"] == code]
            if area_group.empty:
                continue
            ordered_area = _order_stores(area_group)
            title = f"{code} - {area_name}"
            safe_area = area_name.replace("–", "-").replace(" ", "_")
            images[f"Area_{code}_{safe_area}.png"] = _render_map(
                ordered_area,
                title,
                "Peta area lengkap · rute kunjungan dan jarak antar toko",
                area_color,
                tile_cache,
                session,
            )
            for district, district_group in area_group.groupby("Kecamatan", sort=True):
                ordered_district = _order_stores(district_group)
                sector_title = f"{code} - {area_name} - Sektor {district}"
                safe_district = str(district).replace(" ", "_").replace(".", "")
                images[f"Sektor_{code}_{safe_district}.png"] = _render_map(
                    ordered_district,
                    sector_title,
                    "Peta detail sektor kecamatan · rute kunjungan dan jarak antar toko",
                    area_color,
                    tile_cache,
                    session,
                )
    finally:
        session.close()
    return images


def create_screenshot_archive(data):
    images = create_screenshot_images(data)
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as output:
        for filename, image in images.items():
            output.writestr(filename, image)
    return archive.getvalue(), len(images)


def save_screenshots(data, output_dir):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    images = create_screenshot_images(data)
    for filename, image in images.items():
        (output_dir / filename).write_bytes(image)
    archive_path = output_dir / "screenshots_area_dan_sektor.zip"
    with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as output:
        for filename, image in images.items():
            output.writestr(filename, image)
    return archive_path, len(images)


def create_detailed_workbook(data):
    required = {"Nama Toko", "Alamat", "Latitude", "Longitude", "Kode Area", "Area", "Kecamatan"}
    missing = required - set(data.columns)
    if missing:
        raise ValueError(f"Kolom data toko tidak ditemukan: {', '.join(sorted(missing))}.")
    if data.empty:
        raise ValueError("Tidak ada data toko untuk dibuatkan Excel detail.")
    if data[["Latitude", "Longitude"]].isna().any().any():
        raise ValueError("Koordinat Latitude/Longitude kosong; lengkapi data sebelum membuat Excel.")

    groups = []
    for code, (area_name, area_color, _) in AREAS.items():
        area_group = data[data["Kode Area"] == code]
        if area_group.empty:
            continue
        groups.append((f"Area {code} {area_name}", area_group, area_color))
        for district, district_group in area_group.groupby("Kecamatan", sort=True):
            groups.append((f"Sektor {code} {district}", district_group, area_color))

    summary_rows = []
    prepared_groups = []
    for title, group, _ in groups:
        ordered = _order_stores(group)
        _, distances, _ = _route_for(ordered, title)
        area_code = str(ordered.iloc[0]["Kode Area"])
        area_name = str(ordered.iloc[0]["Area"])
        is_area = title.startswith("Area ")
        sector_name = "Semua sektor" if is_area else str(ordered.iloc[0]["Kecamatan"])
        total_distance = sum(distances)
        summary_rows.append({
            "Jenis": "Area" if is_area else "Sektor",
            "Kode Area": area_code,
            "Area": area_name,
            "Sektor / Kecamatan": sector_name,
            "Jumlah Toko": len(ordered),
            "Total Jarak Berkendara (km)": round(total_distance, 3) if distances else None,
            "Keterangan Jarak": (
                "Tidak ada rute: hanya 1 toko" if not distances else "Total jarak rute jalan"
            ),
            "Rute Kunjungan": " → ".join(ordered["Nama Toko"].astype(str)),
        })
        prepared_groups.append((title, ordered, distances))

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        summary = pd.DataFrame(summary_rows)
        summary.to_excel(writer, sheet_name="Ringkasan", index=False)
        for title, ordered, distances in prepared_groups:
            rows = []
            for position, (_, row) in enumerate(ordered.iterrows()):
                distance = distances[position - 1] if position else None
                previous_store = str(ordered.iloc[position - 1]["Nama Toko"]) if position else ""
                rows.append({
                    "No. Urut": position + 1,
                    "Nama Toko": str(row["Nama Toko"]),
                    "Kecamatan": str(row["Kecamatan"]),
                    "Alamat": str(row["Alamat"]),
                    "Latitude": float(row["Latitude"]),
                    "Longitude": float(row["Longitude"]),
                    "Toko Sebelumnya": previous_store,
                    "Jarak (km)": round(distance, 3) if distance is not None else None,
                    "Navigasi Google Maps": (
                        f"https://www.google.com/maps/dir/?api=1"
                        f"&destination={row['Latitude']},{row['Longitude']}"
                    ),
                })
            sheet_name = title[:31]
            pd.DataFrame(rows).to_excel(writer, sheet_name=sheet_name, index=False)

        workbook = writer.book
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo

        for table_number, worksheet in enumerate(workbook.worksheets, start=1):
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.row_dimensions[1].height = 32
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1A73E8")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for column_cells in worksheet.columns:
                column_letter = get_column_letter(column_cells[0].column)
                header = str(column_cells[0].value or "")
                if header in {"Alamat", "Rute Kunjungan"}:
                    width = 48
                elif header == "Keterangan Jarak":
                    width = 34
                elif header == "Navigasi Google Maps":
                    width = 44
                elif header in {"Nama Toko", "Toko Sebelumnya", "Sektor / Kecamatan"}:
                    width = 28
                else:
                    width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 12), 24)
                worksheet.column_dimensions[column_letter].width = width
                if header == "Navigasi Google Maps":
                    for cell in column_cells[1:]:
                        if cell.value:
                            cell.hyperlink = cell.value
                            cell.style = "Hyperlink"
                for cell in column_cells[1:]:
                    cell.alignment = Alignment(vertical="top", wrap_text=header in {"Alamat", "Rute Kunjungan"})
                    if header in {"Jarak (km)", "Total Jarak Berkendara (km)"} and cell.value is not None:
                        cell.number_format = "0.000"
            if worksheet.max_row > 1 and worksheet.max_column > 0:
                table = Table(
                    displayName=f"Data{table_number}",
                    ref=worksheet.dimensions,
                )
                table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                worksheet.add_table(table)

    return buffer.getvalue()


def save_detailed_workbook(data, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = create_detailed_workbook(data)
    output_path.write_bytes(workbook)
    return output_path
